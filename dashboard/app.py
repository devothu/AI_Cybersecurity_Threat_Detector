
from pathlib import Path
import tempfile
import io
import os
import joblib
import numpy as np
import pandas as pd
import streamlit as st


from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


def create_excel_summary_report(results_path, output_path):
    """Create a compact Excel workbook with summary tables and embedded charts."""
    severity_names = ["Critical", "High", "Medium", "Low", "Unscored", "Benign"]
    severity_counts = {name: 0 for name in severity_names}
    total_scanned = 0
    total_attacks = 0
    total_benign = 0

    # Read the scan results in chunks so the report summary does not load
    # the full CSV into memory.
    for batch in pd.read_csv(
        results_path,
        chunksize=10000,
        usecols=["AI_Prediction", "Threat_Severity"],
        low_memory=False,
    ):
        total_scanned += len(batch)
        predictions = batch["AI_Prediction"].astype(str).str.upper()
        total_attacks += int((predictions == "ATTACK").sum())
        total_benign += int((predictions == "BENIGN").sum())
        counts = batch["Threat_Severity"].value_counts()
        for name in severity_names:
            severity_counts[name] += int(counts.get(name, 0))

    wb = Workbook()
    summary = wb.active
    summary.title = "Investigation Summary"
    summary.append(["AI Cybersecurity Threat Detector"])
    summary.append(["Investigation Summary"])
    summary.append(["Metric", "Value"])
    summary.append(["Total records scanned", total_scanned])
    summary.append(["Predicted benign", total_benign])
    summary.append(["Predicted attacks", total_attacks])
    summary.append([
        "Predicted attack share (%)",
        round(total_attacks / total_scanned * 100, 2) if total_scanned else 0,
    ])
    summary.append([])
    summary.append(["Threat Severity", "Records"])
    for name in severity_names:
        summary.append([name, severity_counts[name]])

    # Style headings and column labels.
    summary["A1"].font = Font(bold=True, size=16, color="FFFFFF")
    summary["A1"].fill = PatternFill("solid", fgColor="14213D")
    summary["A2"].font = Font(bold=True, size=13)
    for cell in summary[3]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="264653")
    for cell in summary[9]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="264653")
    summary.column_dimensions["A"].width = 32
    summary.column_dimensions["B"].width = 22
    summary.freeze_panes = "A4"

    # Add a separate charts worksheet with embedded Excel charts.
    charts = wb.create_sheet("Charts")
    charts["A1"] = "Traffic Classification"
    charts["A1"].font = Font(bold=True, size=14)
    charts["A3"] = "Prediction"
    charts["B3"] = "Records"
    charts["A4"] = "BENIGN"
    charts["B4"] = total_benign
    charts["A5"] = "ATTACK"
    charts["B5"] = total_attacks

    traffic_chart = PieChart()
    traffic_chart.title = "Benign vs Attack Predictions"
    traffic_chart.height = 8
    traffic_chart.width = 13
    traffic_chart.add_data(Reference(charts, min_col=2, min_row=3, max_row=5), titles_from_data=True)
    traffic_chart.set_categories(Reference(charts, min_col=1, min_row=4, max_row=5))
    charts.add_chart(traffic_chart, "D2")

    charts["A8"] = "Threat Severity"
    charts["A10"] = "Severity"
    charts["B10"] = "Records"
    for row, name in enumerate(severity_names, start=11):
        charts.cell(row=row, column=1, value=name)
        charts.cell(row=row, column=2, value=severity_counts[name])

    severity_chart = BarChart()
    severity_chart.type = "bar"
    severity_chart.style = 10
    severity_chart.title = "Records by Threat Severity"
    severity_chart.y_axis.title = "Severity"
    severity_chart.x_axis.title = "Records"
    severity_chart.height = 8
    severity_chart.width = 15
    severity_chart.add_data(
        Reference(charts, min_col=2, min_row=10, max_row=10 + len(severity_names)),
        titles_from_data=True,
    )
    severity_chart.set_categories(
        Reference(charts, min_col=1, min_row=11, max_row=10 + len(severity_names))
    )
    charts.add_chart(severity_chart, "D18")

    for sheet in (summary, charts):
        sheet.freeze_panes = "A3"
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="center")
        for col in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(col)].width = max(
                sheet.column_dimensions[get_column_letter(col)].width or 0, 18
            )

    wb.save(output_path)


# ==================================================
# 1. CONFIGURATION
# ==================================================

PROJECT_DIR = Path(__file__).resolve().parent.parent

MODEL_FILE = PROJECT_DIR / "models" / "threat_detector.joblib"
DATA_FILE = PROJECT_DIR / "data" / "processed" / "traffic_cleaned.csv"
REPORT_FILE = PROJECT_DIR / "reports" / "classification_report.txt"
MATRIX_FILE = PROJECT_DIR / "reports" / "confusion_matrix.png"

st.set_page_config(
    page_title="CyberShield AI | Threat Detector",
    page_icon="🛡️",
    layout="wide",
    initial_sidebar_state="expanded",
)


# ==================================================
# 2. CUSTOM DESIGN
# ==================================================

st.markdown("""
<style>
    .stApp {
        background-color: #0b1120;
        color: #e5e7eb;
    }

    [data-testid="stSidebar"] {
        background-color: #111827;
        border-right: 1px solid #263449;
    }

    [data-testid="stMetric"] {
        background-color: #151f31;
        border: 1px solid #29374d;
        padding: 18px;
        border-radius: 12px;
    }

    [data-testid="stMetricLabel"] {
        color: #9ca3af;
    }

    [data-testid="stMetricValue"] {
        color: #f9fafb;
    }

    h1, h2, h3 {
        color: #f9fafb;
    }

    .subtitle {
        color: #9ca3af;
        font-size: 15px;
    }

    .status {
        background-color: #132d2b;
        border: 1px solid #22665c;
        border-radius: 10px;
        padding: 14px;
        color: #a7f3d0;
    }

    .notice {
        background-color: #302719;
        border: 1px solid #806329;
        border-radius: 10px;
        padding: 14px;
        color: #fde68a;
    }

    div[data-testid="stDataFrame"] {
        border: 1px solid #29374d;
        border-radius: 10px;
    }
</style>
""", unsafe_allow_html=True)


# ==================================================
# 3. LOAD PROJECT FILES
# ==================================================

@st.cache_resource
def load_model(path):
    return joblib.load(path)


@st.cache_data(show_spinner="Loading network traffic data...")
def load_dataset(path, modified_time):
    # Load the cleaned dataset so the dashboard can
    # sample across the complete file, not just its first rows.
    return pd.read_csv(path, low_memory=False)


if not MODEL_FILE.exists():
    st.error("Trained model not found. Run src/train_model.py first.")
    st.stop()

if not DATA_FILE.exists():
    st.error("Cleaned dataset not found. Run src/prepare_data.py first.")
    st.stop()

try:
    bundle = load_model(str(MODEL_FILE))
    model = bundle["model"]
    features = bundle["features"]

    df = load_dataset(
        str(DATA_FILE),
        DATA_FILE.stat().st_mtime,
    )

except Exception as error:
    st.error(f"Could not load the project files: {error}")
    st.stop()


if "Label" not in df.columns:
    st.error("The cleaned dataset must contain a Label column.")
    st.stop()

df = df[df["Label"].isin(["BENIGN", "ATTACK"])].copy()

if df.empty:
    st.error("No labelled records are available.")
    st.stop()


# ==================================================
# 4. SIDEBAR NAVIGATION
# ==================================================

with st.sidebar:
    st.markdown("# 🛡️ CyberShield AI")
    st.caption("NETWORK THREAT INTELLIGENCE")
    st.divider()

    page = st.radio(
        "NAVIGATION",
        [
            "Overview",
            "Traffic Explorer",
            "AI Predictions",
            "CSV Threat Scanner",
            "Model Performance",
        ],
    )


    st.divider()

    st.markdown("### Model status")
    st.success("Trained model loaded")

    st.caption("Algorithm: Random Forest")
    st.caption(f"Available records: {len(df):,}")
    st.caption(f"Model features: {len(features)}")

    st.divider()
    st.caption("Portfolio project | Python + Machine Learning")


# ==================================================
# 5. HEADER AND DATA SUMMARY
# ==================================================

st.title("🛡️ AI Cybersecurity Threat Detector")

st.markdown(
    '<p class="subtitle">'
    'Network traffic analytics, machine-learning predictions '
    'and threat investigation.'
    '</p>',
    unsafe_allow_html=True,
)

total_records = len(df)
benign_count = int((df["Label"] == "BENIGN").sum())
attack_count = int((df["Label"] == "ATTACK").sum())

attack_rate = (
    attack_count / total_records * 100
    if total_records else 0
)

# Random sample drawn across the complete dataset.
# This is for visual exploration, not model evaluation.
sample_size = min(20000, total_records)

sample = df.sample(
    n=sample_size,
    random_state=42,
) if total_records else df.copy()


# ==================================================
# 6. OVERVIEW PAGE
# ==================================================

if page == "Overview":

    st.subheader("Security overview")

    c1, c2, c3, c4 = st.columns(4)

    c1.metric("Total records", f"{total_records:,}")
    c2.metric("Benign records", f"{benign_count:,}")
    c3.metric("Attack-labelled records", f"{attack_count:,}")
    c4.metric("Attack share", f"{attack_rate:.2f}%")

    st.markdown("")
    st.markdown(
        '<div class="status">'
        '● MODEL READY &nbsp; | &nbsp; '
        'The trained classifier is available for analysis.'
        '</div>',
        unsafe_allow_html=True,
    )

    st.markdown("")
    left, right = st.columns(2)

    with left:
        st.subheader("Traffic distribution")

        counts = pd.Series(
            {
                "BENIGN": benign_count,
                "ATTACK": attack_count,
            }
        )

        st.bar_chart(counts)

    with right:
        st.subheader("Traffic composition")

        composition = pd.DataFrame(
            {
                "Traffic class": ["BENIGN", "ATTACK"],
                "Records": [benign_count, attack_count],
            }
        )

        st.dataframe(
            composition,
            hide_index=True,
            width="stretch",
        )

        st.caption(
            "These figures describe the full cleaned dataset, "
            "not live network activity."
        )

    st.subheader("Sample network records")

    display_columns = [
        column for column in [
            "Label",
            "Protocol",
            "Flow Duration",
            "Total Fwd Packets",
            "Total Backward Packets",
            "Flow Bytes/s",
        ]
        if column in sample.columns
    ]

    st.dataframe(
        sample[display_columns].head(10),
        hide_index=True,
        width="stretch",
    )

    st.markdown(
        '<div class="notice">'
        '<strong>Investigation note:</strong> '
        'Dataset labels and model predictions are not proof of a '
        'real-world attack. This application analyses a static dataset.'
        '</div>',
        unsafe_allow_html=True,
    )


# ==================================================
# 7. TRAFFIC EXPLORER
# ==================================================

elif page == "Traffic Explorer":

    st.subheader("Network traffic explorer")
    st.write(
        "Filter the dataset to inspect records from either "
        "traffic class."
    )

    selected_label = st.selectbox(
        "Filter by reference label",
        ["ALL", "BENIGN", "ATTACK"],
    )

    filtered = (
        sample
        if selected_label == "ALL"
        else sample[sample["Label"] == selected_label]
    )

    st.metric("Records in current view", f"{len(filtered):,}")

    search = st.text_input(
        "Search column names",
        placeholder="Example: Flow Duration",
    )

    available_columns = list(filtered.columns)

    if search.strip():
        matching = [
            column for column in available_columns
            if search.lower() in column.lower()
        ]

        if "Label" in available_columns and "Label" not in matching:
            matching.insert(0, "Label")

        available_columns = matching

    selected_columns = st.multiselect(
        "Columns to display",
        options=available_columns,
        default=available_columns[:8],
    )

    if selected_columns:
        st.dataframe(
            filtered[selected_columns],
            hide_index=True,
            width="stretch",
        )
    else:
        st.info("Select at least one column to display.")

    st.caption(
        "The explorer displays a reproducible random sample of up to "
        "20,000 records from the full dataset."
    )


# ==================================================
# 8. AI PREDICTIONS
# ==================================================

elif page == "AI Predictions":

    st.subheader("Inspect a network traffic record")

    st.write(
        "Choose a record from the sample. The model will classify "
        "its numerical network-flow features."
    )

    record_indices = sample.index.tolist()

    selected_index = st.selectbox(
        "Select record",
        options=record_indices,
        format_func=lambda index: f"Record {index}",
    )

    record = df.loc[[selected_index]]

    X_record = record.reindex(columns=features)
    X_record = X_record.apply(pd.to_numeric, errors="coerce")
    X_record = X_record.replace([np.inf, -np.inf], np.nan)

    try:
        prediction = model.predict(X_record)[0]

        left, right = st.columns(2)

        with left:
            st.markdown("### AI classification")

            if prediction == "ATTACK":
                st.error("⚠️ Potentially suspicious")
            else:
                st.success("✓ Classified as BENIGN")

            st.metric("Predicted class", str(prediction))

        with right:
            st.markdown("### Dataset reference")

            actual_label = str(record["Label"].iloc[0])
            st.metric("Reference label", actual_label)

            if prediction == actual_label:
                st.info("Prediction matches the dataset label.")
            else:
                st.warning(
                    "Prediction differs from the dataset label. "
                    "Investigate this record."
                )

        if hasattr(model, "predict_proba"):
            probabilities = model.predict_proba(X_record)[0]

            scores = pd.DataFrame(
                {
                    "Class": model.classes_,
                    "Model score": probabilities,
                }
            )

            st.subheader("Model scores")
            st.bar_chart(scores.set_index("Class"))

            st.caption(
                "Scores are model outputs and may not be calibrated "
                "probabilities of a real-world attack."
            )

        with st.expander("Inspect numerical traffic features"):
            st.dataframe(
                X_record.T.rename(columns={selected_index: "Value"}),
                width="stretch",
            )

    except Exception as error:
        st.error(f"Could not classify this record: {error}")


# ==================================================
# CSV THREAT SCANNER
# ==================================================




elif page == "CSV Threat Scanner":
    from pathlib import Path

    st.title("CSV Threat Scanner")
    st.write("Scan network traffic in batches and investigate suspicious records.")

    reports_dir = PROJECT_DIR / "reports"
    reports_dir.mkdir(parents=True, exist_ok=True)
    results_path = reports_dir / "cybersecurity_scan_results.csv"

    uploaded_file = st.file_uploader(
        "Upload a network traffic CSV file",
        type=["csv"],
        key="batch_csv_upload"
    )

    batch_size = st.selectbox(
        "Records processed per batch",
        [1000, 5000, 10000],
        index=2
    )

    if uploaded_file is not None:
        st.caption(
            f"File size: {uploaded_file.size / (1024 * 1024):.2f} MB"
        )

        if st.button("Start Threat Scan", type="primary"):
            temp_path = None

            try:
                # Save the upload to disk so it can be read in batches.
                fd, temp_path = tempfile.mkstemp(suffix=".csv")
                os.close(fd)

                with open(temp_path, "wb") as input_file:
                    input_file.write(uploaded_file.getvalue())

                # Write to a temporary results file first.
                fd, temp_results_path = tempfile.mkstemp(suffix=".csv")
                os.close(fd)

                total_records = 0
                benign_count = 0
                attack_count = 0
                first_batch = True

                progress = st.progress(0)
                status = st.empty()

                # Count records for progress reporting.
                total_expected = sum(
                    len(batch)
                    for batch in pd.read_csv(
                        temp_path,
                        chunksize=batch_size,
                        low_memory=False
                    )
                )

                if total_expected == 0:
                    raise ValueError("The uploaded CSV contains no records.")

                with open(
                    temp_results_path,
                    "w",
                    encoding="utf-8",
                    newline=""
                ) as output_file:

                    for chunk in pd.read_csv(
                        temp_path,
                        chunksize=batch_size,
                        low_memory=False
                    ):
                        chunk.columns = chunk.columns.str.strip()

                        if chunk.columns.duplicated().any():
                            raise ValueError(
                                "The CSV contains duplicate column names."
                            )

                        missing_features = [
                            feature for feature in features
                            if feature not in chunk.columns
                        ]

                        if missing_features:
                            raise ValueError(
                                "Missing required model features: "
                                + ", ".join(missing_features[:20])
                            )

                        # Match the exact feature order used in training.
                        X_batch = chunk[features].copy()

                        for column in X_batch.columns:
                            X_batch[column] = pd.to_numeric(
                                X_batch[column], errors="coerce"
                            )

                        X_batch = X_batch.replace(
                            [np.inf, -np.inf], np.nan
                        )

                        # This model must support the resulting missing
                        # values, or preprocessing must match training.
                        predictions = model.predict(X_batch)
                        chunk["AI_Prediction"] = predictions

                        # Calculate attack probability if supported.
                        attack_probability = np.full(
                            len(chunk), np.nan, dtype=float
                        )

                        if hasattr(model, "predict_proba"):
                            probabilities = model.predict_proba(X_batch)
                            classes = [
                                str(label).strip().upper()
                                for label in model.classes_
                            ]

                            attack_indices = [
                                i for i, label in enumerate(classes)
                                if label == "ATTACK"
                            ]

                            if attack_indices:
                                attack_probability = (
                                    probabilities[:, attack_indices].sum(axis=1)
                                    * 100
                                )

                        chunk["Attack_Probability"] = np.round(
                            attack_probability, 2
                        )

                        # Assign preliminary severity to predictions.
                        severity = []

                        for prediction, probability in zip(
                            predictions, attack_probability
                        ):
                            if str(prediction).strip().upper() != "ATTACK":
                                severity.append("Benign")
                            elif np.isnan(probability):
                                severity.append("Unscored")
                            elif probability >= 90:
                                severity.append("Critical")
                            elif probability >= 75:
                                severity.append("High")
                            elif probability >= 50:
                                severity.append("Medium")
                            else:
                                severity.append("Low")

                        chunk["Threat_Severity"] = severity

                        benign_count += sum(
                            str(value).strip().upper() == "BENIGN"
                            for value in predictions
                        )
                        attack_count += sum(
                            str(value).strip().upper() == "ATTACK"
                            for value in predictions
                        )

                        chunk.to_csv(
                            output_file,
                            index=False,
                            header=first_batch
                        )
                        first_batch = False

                        total_records += len(chunk)
                        progress.progress(
                            min(total_records / total_expected, 1.0)
                        )
                        status.write(
                            f"Processed {total_records:,} "
                            f"of {total_expected:,} records..."
                        )

                # Only replace the previous report after a successful scan.
                os.replace(temp_results_path, results_path)

                st.success("Threat scan completed successfully!")

                col1, col2, col3 = st.columns(3)
                col1.metric("Records scanned", f"{total_records:,}")
                col2.metric("Predicted benign", f"{benign_count:,}")
                col3.metric("Predicted attacks", f"{attack_count:,}")

                st.metric(
                    "Predicted attack share",
                    f"{attack_count / total_records * 100:.2f}%"
                    if total_records else "0%"
                )

                st.caption(
                    "Severity is based on model scores, not confirmed "
                    "real-world impact."
                )

            except Exception as error:
                st.error(f"Scan failed: {error}")

            finally:
                if temp_path and os.path.exists(temp_path):
                    os.remove(temp_path)

                if (
                    "temp_results_path" in locals()
                    and os.path.exists(temp_results_path)
                ):
                    os.remove(temp_results_path)

    # Investigation tools are shown only on this page.
    st.divider()
    st.subheader("Suspicious Traffic Investigation")

    if results_path.exists():
        severity_options = [
            "All", "Critical", "High", "Medium",
            "Low", "Unscored", "Benign"
        ]

        selected_severity = st.selectbox(
            "Filter by severity",
            severity_options
        )

        selected_prediction = st.selectbox(
            "Filter by AI prediction",
            ["All", "ATTACK", "BENIGN"]
        )

        minimum_probability = st.slider(
            "Minimum attack probability (%)",
            min_value=0,
            max_value=100,
            value=50
        )

        if st.button("Apply Investigation Filters"):
            matching_rows = []
            matched_count = 0

            for batch in pd.read_csv(
                results_path,
                chunksize=10000,
                low_memory=False
            ):
                if selected_severity != "All":
                    batch = batch[
                        batch["Threat_Severity"] == selected_severity
                    ]

                if selected_prediction != "All":
                    batch = batch[
                        batch["AI_Prediction"].astype(str).str.upper()
                        == selected_prediction
                    ]

                if "Attack_Probability" in batch.columns:
                    scores = pd.to_numeric(
                        batch["Attack_Probability"], errors="coerce"
                    )

                    # Benign records have no need to pass an attack
                    # probability threshold when severity is "Benign".
                    if selected_severity != "Benign":
                        batch = batch[
                            scores.loc[batch.index] >= minimum_probability
                        ]

                matched_count += len(batch)

                remaining = 500 - len(matching_rows)
                if remaining > 0:
                    matching_rows.extend(
                        batch.head(remaining).to_dict("records")
                    )

            st.metric("Matching records", f"{matched_count:,}")
            st.caption("Showing a maximum of 500 matching records.")

            if matching_rows:
                st.dataframe(
                    pd.DataFrame(matching_rows),
                    hide_index=True,
                    width="stretch"
                )
            else:
                st.info("No records match these filters.")

        st.download_button(
            "Download Full Scan Results (CSV)",
            data=results_path.read_bytes(),
            file_name="cybersecurity_scan_results.csv",
            mime="text/csv",
            width="stretch"
        )

        # Build an Excel workbook containing the investigation summary and
        # embedded charts. The full row-level results remain available above
        # as CSV to avoid creating an extremely large Excel workbook.
        excel_report_path = reports_dir / "cybersecurity_investigation_report.xlsx"
        try:
            create_excel_summary_report(results_path, excel_report_path)
            st.download_button(
                "Download Excel Summary & Charts",
                data=excel_report_path.read_bytes(),
                file_name="cybersecurity_investigation_report.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                width="stretch",
            )
        except Exception as error:
            st.error(f"Could not create the Excel report: {error}")

        st.subheader("Investigation Summary")

        severity_counts = {
            "Critical": 0,
            "High": 0,
            "Medium": 0,
            "Low": 0,
            "Unscored": 0,
            "Benign": 0
        }

        total_scanned = 0
        total_attacks = 0

        for batch in pd.read_csv(
            results_path,
            chunksize=10000,
            usecols=["AI_Prediction", "Threat_Severity"],
            low_memory=False
        ):
            total_scanned += len(batch)

            total_attacks += int(
                (
                    batch["AI_Prediction"].astype(str).str.upper()
                    == "ATTACK"
                ).sum()
            )

            counts = batch["Threat_Severity"].value_counts()

            for name in severity_counts:
                severity_counts[name] += int(counts.get(name, 0))

        summary_df = pd.DataFrame({
            "Severity": list(severity_counts.keys()),
            "Records": list(severity_counts.values())
        })

        a, b = st.columns(2)
        a.metric("Total scanned", f"{total_scanned:,}")
        b.metric("Predicted attacks", f"{total_attacks:,}")

        st.bar_chart(summary_df.set_index("Severity"))
        st.dataframe(summary_df, hide_index=True, width="stretch")

        st.download_button(
            "Download Investigation Summary",
            data=summary_df.to_csv(index=False).encode("utf-8"),
            file_name="cybersecurity_investigation_summary.csv",
            mime="text/csv",
            width="stretch"
        )

    else:
        st.info("Run a successful scan to create investigation results.")


# ==================================================
# 10. FOOTER
# ==================================================

st.divider()

st.caption(
    "CyberShield AI | Educational cybersecurity analytics project. "
    "Predictions are investigative indicators, not definitive proof "
    "of malicious activity."
)
